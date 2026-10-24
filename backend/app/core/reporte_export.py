"""Exportar cualquier reporte (CU07 de cada tienda y CU20 de la plataforma):
Imprimir/PDF, Excel, CSV y HTML, para descargar o enviar por correo.

Un reporte es una tabla: titulo, subtitulo (filtros/periodo), columnas con
su tipo (texto, numero, moneda, fecha), filas y, si hay, una fila de totales.
"""

import csv
import html
import io
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from typing import Any

from fastapi import HTTPException, Response, status
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.email import send_email

COLOR = "1F2A44"
FORMATOS = {
    "pdf": ("application/pdf", "pdf"),
    "excel": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "csv": ("text/csv; charset=utf-8", "csv"),
    "html": ("text/html; charset=utf-8", "html"),
}
MAX_FILAS = 5000


@dataclass
class Columna:
    clave: str
    etiqueta: str
    tipo: str = "texto"  # texto | numero | moneda | fecha


@dataclass
class Reporte:
    titulo: str
    columnas: list[Columna]
    filas: list[dict[str, Any]]
    subtitulo: str | None = None
    totales: dict[str, Any] | None = None
    marca: str = "Arlani Boutiques"

    @property
    def generado(self) -> str:
        # Hora de Bolivia (UTC-4).
        return (datetime.now(timezone.utc) - timedelta(hours=4)).strftime("%d/%m/%Y %H:%M")


def desde_listas(titulo: str, headers: list[str], filas: list[list[Any]], **kw) -> Reporte:
    """Para los reportes que ya armaban encabezados y filas como listas."""
    columnas = []
    for i, h in enumerate(headers):
        valores = [f[i] for f in filas if f[i] not in (None, "")]
        numerico = bool(valores) and all(isinstance(v, (int, float, Decimal)) and not isinstance(v, bool) for v in valores)
        tipo = ("moneda" if "Bs" in h or "Total" in h or "Monto" in h else "numero") if numerico else "texto"
        columnas.append(Columna(f"c{i}", h, tipo))
    return Reporte(titulo, columnas, [{f"c{i}": v for i, v in enumerate(f)} for f in filas], **kw)


def _texto(valor: Any, tipo: str) -> str:
    if valor is None or valor == "":
        return ""
    if tipo == "moneda":
        return f"{Decimal(str(valor)):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if tipo == "numero" and isinstance(valor, (float, Decimal)) and Decimal(str(valor)) % 1:
        return f"{Decimal(str(valor)):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M")
    return str(valor)


def _numero(valor: Any) -> Any:
    if isinstance(valor, Decimal):
        return float(valor)
    return valor


# ------------------------------------------------------------- formatos


def _csv(r: Reporte) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c.etiqueta for c in r.columnas])
    for f in r.filas:
        w.writerow([_numero(f.get(c.clave)) if c.tipo in ("numero", "moneda") else _texto(f.get(c.clave), c.tipo) for c in r.columnas])
    if r.totales:
        w.writerow([_numero(r.totales.get(c.clave, "Total" if i == 0 else "")) for i, c in enumerate(r.columnas)])
    # BOM: Excel abre bien las tildes del CSV.
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _excel(r: Reporte) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    n = max(len(r.columnas), 1)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws.cell(row=1, column=1, value=f"{r.marca} — {r.titulo}").font = Font(bold=True, size=14, color=COLOR)
    ws.cell(row=2, column=1, value=(r.subtitulo + " · " if r.subtitulo else "") + f"Generado el {r.generado}").font = Font(
        italic=True, color="666666"
    )
    fila = 4
    for i, c in enumerate(r.columnas, start=1):
        celda = ws.cell(row=fila, column=i, value=c.etiqueta)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(start_color=COLOR, end_color=COLOR, fill_type="solid")
        celda.alignment = Alignment(horizontal="right" if c.tipo in ("numero", "moneda") else "left")
    for f in r.filas:
        fila += 1
        for i, c in enumerate(r.columnas, start=1):
            v = f.get(c.clave)
            celda = ws.cell(row=fila, column=i, value=_numero(v) if c.tipo in ("numero", "moneda") else _texto(v, c.tipo))
            if c.tipo == "moneda":
                celda.number_format = "#,##0.00"
    if r.totales:
        fila += 1
        for i, c in enumerate(r.columnas, start=1):
            v = r.totales.get(c.clave, "Total" if i == 1 else None)
            celda = ws.cell(row=fila, column=i, value=_numero(v))
            celda.font = Font(bold=True)
            if c.tipo == "moneda":
                celda.number_format = "#,##0.00"
    for i, c in enumerate(r.columnas, start=1):
        largo = max([len(c.etiqueta)] + [len(_texto(f.get(c.clave), c.tipo)) for f in r.filas[:500]] + [10])
        ws.column_dimensions[get_column_letter(i)].width = min(largo + 2, 45)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def html_tabla(r: Reporte) -> str:
    """La tabla con estilos en linea (sirve suelta y dentro de un correo)."""
    e = html.escape
    th = "".join(
        f'<th style="background:#{COLOR};color:#fff;padding:8px 10px;font-size:12px;'
        f'text-align:{"right" if c.tipo in ("numero", "moneda") else "left"}">{e(c.etiqueta)}</th>'
        for c in r.columnas
    )
    filas = []
    for k, f in enumerate(r.filas):
        fondo = "#f7f7f5" if k % 2 else "#ffffff"
        tds = "".join(
            f'<td style="padding:7px 10px;border-bottom:1px solid #eee;font-size:13px;'
            f'text-align:{"right" if c.tipo in ("numero", "moneda") else "left"}">{e(_texto(f.get(c.clave), c.tipo))}</td>'
            for c in r.columnas
        )
        filas.append(f'<tr style="background:{fondo}">{tds}</tr>')
    if r.totales:
        tds = "".join(
            f'<td style="padding:8px 10px;font-weight:700;font-size:13px;border-top:2px solid #{COLOR};'
            f'text-align:{"right" if c.tipo in ("numero", "moneda") else "left"}">'
            f'{e(_texto(r.totales.get(c.clave, "Total" if i == 0 else ""), c.tipo))}</td>'
            for i, c in enumerate(r.columnas)
        )
        filas.append(f"<tr>{tds}</tr>")
    return f'<table style="border-collapse:collapse;width:100%;font-family:Arial,sans-serif"><thead><tr>{th}</tr></thead><tbody>{"".join(filas)}</tbody></table>'


def _html(r: Reporte) -> bytes:
    e = html.escape
    sub = f"<p style='color:#666;margin:0 0 16px'>{e(r.subtitulo)}</p>" if r.subtitulo else ""
    doc = f"""<!doctype html><html lang="es"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><title>{e(r.titulo)} — {e(r.marca)}</title></head>
<body style="margin:0;padding:24px;background:#fff;color:#222;font-family:Arial,sans-serif">
<p style="margin:0;color:#888;font-size:12px;text-transform:uppercase;letter-spacing:.06em">{e(r.marca)}</p>
<h1 style="margin:4px 0 6px;color:#{COLOR};font-size:22px">{e(r.titulo)}</h1>{sub}
{html_tabla(r)}
<p style="color:#999;font-size:11px;margin-top:16px">Generado el {r.generado} · {len(r.filas)} filas</p>
</body></html>"""
    return doc.encode("utf-8")


def _latin(texto: str) -> str:
    # Las fuentes base de PDF son latin-1: tildes y ñ sí; otros símbolos se reemplazan.
    return texto.replace("—", "-").replace("·", "-").encode("latin-1", "replace").decode("latin-1")


def _pdf(r: Reporte) -> bytes:
    horizontal = len(r.columnas) > 5
    pdf = FPDF(orientation="L" if horizontal else "P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    ancho = pdf.w - 2 * pdf.l_margin
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(130, 130, 130)
    pdf.cell(0, 5, _latin(r.marca.upper()), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "B", 16)
    pdf.set_text_color(31, 42, 68)
    pdf.multi_cell(0, 8, _latin(r.titulo), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(100, 100, 100)
    pdf.multi_cell(0, 5, _latin((r.subtitulo + " - " if r.subtitulo else "") + f"Generado el {r.generado}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    # Anchos segun el contenido (con minimo y maximo).
    largos = [
        max([len(c.etiqueta)] + [len(_texto(f.get(c.clave), c.tipo)) for f in r.filas[:300]]) for c in r.columnas
    ]
    largos = [min(max(x, 6), 40) for x in largos]
    anchos = [ancho * x / sum(largos) for x in largos] if largos else []

    def encabezado() -> None:
        pdf.set_font("Helvetica", "B", 8)
        pdf.set_fill_color(31, 42, 68)
        pdf.set_text_color(255, 255, 255)
        for c, w in zip(r.columnas, anchos):
            pdf.cell(w, 7, _latin(c.etiqueta)[: int(w / 1.6)], border=0, fill=True,
                     align="R" if c.tipo in ("numero", "moneda") else "L")
        pdf.ln()
        pdf.set_font("Helvetica", "", 8)
        pdf.set_text_color(40, 40, 40)

    encabezado()
    for k, f in enumerate(r.filas):
        if pdf.get_y() > pdf.h - 18:
            pdf.add_page()
            encabezado()
        pdf.set_fill_color(247, 247, 245) if k % 2 else pdf.set_fill_color(255, 255, 255)
        for c, w in zip(r.columnas, anchos):
            pdf.cell(w, 6, _latin(_texto(f.get(c.clave), c.tipo))[: int(w / 1.5)], border=0, fill=True,
                     align="R" if c.tipo in ("numero", "moneda") else "L")
        pdf.ln()
    if r.totales:
        pdf.set_font("Helvetica", "B", 8)
        pdf.set_draw_color(31, 42, 68)
        pdf.line(pdf.l_margin, pdf.get_y(), pdf.l_margin + ancho, pdf.get_y())
        for i, (c, w) in enumerate(zip(r.columnas, anchos)):
            v = r.totales.get(c.clave, "Total" if i == 0 else "")
            pdf.cell(w, 7, _latin(_texto(v, c.tipo)), align="R" if c.tipo in ("numero", "moneda") else "L")
        pdf.ln()
    return bytes(pdf.output())


def generar(r: Reporte, formato: str) -> tuple[bytes, str, str]:
    if formato not in FORMATOS:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Formato no válido.")
    if len(r.filas) > MAX_FILAS:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"El reporte tiene más de {MAX_FILAS} filas: acota los filtros.")
    contenido = {"pdf": _pdf, "excel": _excel, "csv": _csv, "html": _html}[formato](r)
    mime, ext = FORMATOS[formato]
    return contenido, mime, ext


def _nombre(r: Reporte, ext: str) -> str:
    base = "".join(ch if ch.isalnum() else "_" for ch in r.titulo.lower()).strip("_")[:50] or "reporte"
    return f"{base}.{ext}"


def descargar(r: Reporte, formato: str) -> Response:
    contenido, mime, ext = generar(r, formato)
    return Response(content=contenido, media_type=mime, headers={"Content-Disposition": f'attachment; filename="{_nombre(r, ext)}"'})


def enviar(r: Reporte, formato: str, correo: str) -> dict:
    """Manda el reporte como archivo adjunto (en HTML, ademas en el cuerpo)."""
    contenido, mime, ext = generar(r, formato)
    e = html.escape
    cuerpo = f"""<div style="font-family:Arial,sans-serif;color:#222">
<p style="margin:0;color:#888;font-size:12px;text-transform:uppercase">{e(r.marca)}</p>
<h2 style="margin:4px 0;color:#{COLOR}">{e(r.titulo)}</h2>
<p style="color:#666;margin:0 0 14px">{e(r.subtitulo or "")}</p>
<p>Te enviamos el reporte en formato <strong>{ext.upper()}</strong> como archivo adjunto ({len(r.filas)} filas, generado el {r.generado}).</p>
{html_tabla(r) if formato == "html" else ""}
</div>"""
    try:
        send_email(correo, f"Reporte: {r.titulo} — {r.marca}", cuerpo, adjuntos=[(_nombre(r, ext), contenido, mime.split(";")[0])])
    except Exception as exc:  # noqa: BLE001 - se informa al usuario sin cortar la app
        raise HTTPException(status.HTTP_502_BAD_GATEWAY, f"No se pudo enviar el correo: {str(exc)[:160]}") from exc
    return {"enviado": True, "correo": correo, "formato": ext}
